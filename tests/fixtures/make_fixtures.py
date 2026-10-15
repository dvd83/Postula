"""Fabrique des CV d'essai avec photo (personne fictive) pour les tests de récupération de photo.
python3 tests/fixtures/make_fixtures.py   (nécessite Pillow, python-docx, reportlab et LibreOffice)"""
import os, subprocess
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Mm, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

D = os.path.dirname(os.path.abspath(__file__))
# Portrait fictif : fond bleu, visage et épaules stylisés
img = Image.new("RGB", (400, 500), (60, 110, 170)); g = ImageDraw.Draw(img)
g.ellipse((120, 80, 280, 270), fill=(236, 196, 160)); g.rectangle((90, 300, 310, 500), fill=(40, 40, 50))
g.ellipse((60, 280, 340, 520), fill=(40, 40, 50)); g.ellipse((165, 150, 185, 170), fill=(30, 30, 30)); g.ellipse((215, 150, 235, 170), fill=(30, 30, 30))
photo = os.path.join(D, "portrait.jpg"); img.save(photo, quality=85)
# Petit logo (doit être ignoré : trop petit)
logo = os.path.join(D, "logo.png"); Image.new("RGB", (60, 60), (200, 30, 30)).save(logo)

LINES = ["Camille Rochat", "Chef de projet digital", "camille.rochat@example.ch · +41 79 123 45 67 · Lausanne",
         "EXPÉRIENCE", "Chef de projet digital – Le Temps, Genève (2021 – aujourd'hui)",
         "Pilotage de la refonte du site, animation Scrum, suivi SEO, budget annuel de 480 000 CHF.",
         "Product owner – Manor, Bâle (2018 – 2021)", "Backlog e-commerce, click & collect dans 60 magasins.",
         "FORMATION", "Master en systèmes d'information – HEC Lausanne (2016)", "COMPÉTENCES", "Agile, Scrum, Jira, SEO, Google Analytics"]

# 1) Word avec photo (tableau : texte à gauche, photo à droite) + logo en en-tête
doc = Document()
doc.sections[0].header.paragraphs[0].add_run().add_picture(logo, width=Mm(8))
t = doc.add_table(rows=1, cols=2)
t.cell(0, 0).paragraphs[0].add_run(LINES[0]).bold = True
t.cell(0, 0).add_paragraph(LINES[1]); t.cell(0, 0).add_paragraph(LINES[2])
t.cell(0, 1).paragraphs[0].add_run().add_picture(photo, width=Mm(32))
for l in LINES[3:]:
    p = doc.add_paragraph(l)
    if l.isupper():
        p.runs[0].bold = True
docx_path = os.path.join(D, "cv_photo.docx"); doc.save(docx_path)

# 2) Le même exporté en PDF par LibreOffice (comme « Enregistrer en PDF » depuis Word)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", D, docx_path], check=True, capture_output=True)
os.replace(os.path.join(D, "cv_photo.pdf"), os.path.join(D, "cv_photo_word.pdf"))

# 3) PDF « design » : photo découpée en rond, posée dans un bloc de couleur
c = canvas.Canvas(os.path.join(D, "cv_photo_rond.pdf"), pagesize=A4)
W, H = A4
c.setFillColorRGB(0.1, 0.13, 0.2); c.rect(0, H - 60 * mm, W, 60 * mm, fill=1, stroke=0)
c.saveState(); p = c.beginPath(); p.circle(W - 40 * mm, H - 30 * mm, 20 * mm); c.clipPath(p, stroke=0, fill=0)
c.drawImage(photo, W - 60 * mm, H - 55 * mm, 40 * mm, 50 * mm); c.restoreState()
c.drawImage(logo, 15 * mm, 15 * mm, 6 * mm, 6 * mm)
c.setFillColorRGB(1, 1, 1); c.setFont("Helvetica-Bold", 22); c.drawString(20 * mm, H - 28 * mm, "Camille Rochat")
c.setFont("Helvetica", 11); c.drawString(20 * mm, H - 36 * mm, "Chef de projet digital")
c.setFillColorRGB(0, 0, 0); y = H - 75 * mm
for l in LINES[2:]:
    c.setFont("Helvetica-Bold" if l.isupper() else "Helvetica", 10); c.drawString(20 * mm, y, l); y -= 7 * mm
c.save()
# 4) PDF façon modèle Word : grands aplats/dégradés de couleur enregistrés comme images, plus gros que la photo
blk1 = Image.new("RGB", (440, 754)); g1 = ImageDraw.Draw(blk1)
for y in range(754):
    g1.line([(0, y), (440, y)], fill=(70, 140 + y // 40, 190))
blk2 = Image.new("RGB", (425, 580), (75, 145, 195))
b1, b2 = os.path.join(D, "_blk1.png"), os.path.join(D, "_blk2.png"); blk1.save(b1); blk2.save(b2)
c = canvas.Canvas(os.path.join(D, "cv_photo_aplats.pdf"), pagesize=A4)
c.drawImage(b1, 135 * mm, 20 * mm, 60 * mm, 103 * mm); c.drawImage(b2, 135 * mm, 130 * mm, 58 * mm, 79 * mm)
c.drawImage(photo, 150 * mm, H - 48 * mm, 38 * mm, 40 * mm)
c.setFont("Helvetica-Bold", 22); c.drawString(20 * mm, H - 28 * mm, "Camille Rochat"); y = H - 60 * mm
for l in LINES[1:]:
    c.setFont("Helvetica-Bold" if l.isupper() else "Helvetica", 10); c.drawString(20 * mm, y, l); y -= 7 * mm
c.save(); os.remove(b1); os.remove(b2)
print("ok")
